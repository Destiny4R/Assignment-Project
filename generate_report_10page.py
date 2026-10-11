"""Condensed <=10-page report. Overwrites the same .docx."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

OUT = r"C:\Users\BDIC-004\Documents\Visual Studio 2026\Projects\AssignmentProject\GIS_Early_Warning_Conflict_Monitoring_Report_10Page.docx"
doc = Document()

st = doc.styles['Normal']
st.font.name = 'Calibri'
st.font.size = Pt(10)
st.paragraph_format.space_after = Pt(3)
st.paragraph_format.space_before = Pt(0)
st.paragraph_format.line_spacing = 1.0

for i, sz in [(1, 14), (2, 11.5), (3, 10.5)]:
    hs = doc.styles[f'Heading {i}']
    hs.font.name = 'Calibri'
    hs.font.size = Pt(sz)
    hs.font.color.rgb = RGBColor(0x1F, 0x3B, 0x63)
    hs.paragraph_format.space_after = Pt(3)
    hs.paragraph_format.space_before = Pt(6 if i > 1 else 8)

for s in doc.sections:
    s.top_margin = Inches(0.6)
    s.bottom_margin = Inches(0.6)
    s.left_margin = Inches(0.7)
    s.right_margin = Inches(0.7)

ft = doc.sections[0].footer.paragraphs[0]
ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = ft.add_run()
f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), 'PAGE')
rr = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = '1'; rr.append(t); f.append(rr)
r._r.append(f)
ft.add_run('  |  GIS Early Warning — Project Report (10-page edition)').font.size = Pt(7)
for x in ft.runs:
    x.font.size = Pt(7); x.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

def P(text, bold=False, italic=False, size=None, align=None, after=None):
    p = doc.add_paragraph()
    rn = p.add_run(text); rn.bold = bold; rn.italic = italic
    if size: rn.font.size = Pt(size)
    if align: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p

def B(text, pre=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(0.3)
    if pre:
        x = p.add_run(pre); x.bold = True; x.font.size = Pt(10)
    y = p.add_run(text); y.font.size = Pt(10)
    return p

def C(code, title=None):
    if title:
        t = doc.add_paragraph(); x = t.add_run(title)
        x.bold = True; x.font.size = Pt(8.5); x.font.color.rgb = RGBColor(0x1F, 0x3B, 0x63)
        t.paragraph_format.space_after = Pt(1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), 'F2F2F2'); pPr.append(shd)
    x = p.add_run(code); x.font.name = 'Consolas'; x.font.size = Pt(7.5)
    return p

def T(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ''
        x = c.paragraphs[0].add_run(h); x.bold = True; x.font.size = Pt(8.5)
        s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:fill'), 'D9E2F3')
        c._tc.get_or_add_tcPr().append(s)
    for row in rows:
        cs = t.add_row().cells
        for i, v in enumerate(row):
            cs[i].text = ''
            x = cs[i].paragraphs[0].add_run(str(v)); x.font.size = Pt(8.5)
            cs[i].paragraphs[0].paragraph_format.space_after = Pt(1)
    d = doc.add_paragraph(); d.paragraph_format.space_after = Pt(2); d.paragraph_format.space_before = Pt(0)
    r_ = d.add_run(''); r_.font.size = Pt(2)
    return t

def H1(x): return doc.add_heading(x, level=1)
def H2(x): return doc.add_heading(x, level=2)

# ---------- COVER ----------
P('GIS-Based Early Warning and Real-Time Conflict Monitoring System', bold=True, size=20, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
P('Farmer–Herder Conflict (Benue State Focus) — Project Development Report', size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
P('ASP.NET Core 8 · Razor Pages · EF Core 8 · MySQL (Pomelo) · Identity · SignalR · Leaflet.js / OSM', size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
T(['Item', 'Detail'], [
    ['Solution', 'AssignmentProject.slnx → Assignment/Assignment.csproj (net8.0)'],
    ['DB / Seed', 'FarmerHerderConflictDb (MySQL localhost:3306); admin@test.com / Admin123!'],
    ['Run / Publish', 'http://localhost:5278 · https://localhost:7006 · FolderProfile → Desktop\\MyAssignment'],
    ['Evidence', 'Migration 20261008171902_InitialCreate (665 lines); Logs/app-20261008.log (3,027 lines); 08 Oct 2026'],
])
P('Condensed 10-page edition. Written from the as-built tree: Program.cs, appsettings.json, Models/Data/DTOs/Services/Hubs, all Razor Pages + wwwroot/js, migration, log, launch/publish profiles, and Implementation_Plan.md (cited as “Plan §n”).', italic=True, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
P(f'Generated {datetime.date.today().isoformat()} · C:\\Users\\BDIC-004\\Documents\\Visual Studio 2026\\Projects\\AssignmentProject', size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER)

# ---------- CONTENTS ----------
H1('Contents')
for title, pg in [('Abstract','2'),('1  Introduction','2'),('2  Requirements and Threat Model','3'),('3  System Design','4'),('4  Implementation','5'),('5  Testing and Evaluation','7'),('6  Deployment and Operating Policy','7'),('7  Security Considerations','8'),('8  Limitations and Future Work','9'),('9  Conclusion','9'),('References','10'),('Appendix A: Command Reference','10'),('Appendix B: Source Code Layout','10')]:
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0)
    a = p.add_run(title + f'  .....  {pg}'); a.font.size = Pt(9.5)

# ---------- ABSTRACT ----------
H1('Abstract')
doc.add_paragraph('This report documents a GIS-based early warning and real-time monitoring system for farmer–herder conflicts, built as an ASP.NET Core 8 Razor Pages app on MySQL. It links (i) anonymous consented GPS/incident reporting (browser geolocation + Nominatim reverse-geocode), (ii) an officer-only live loop (dashboard stats, Leaflet live map, triage worklist, GeoJSON operations map with hotspot overlays, pushed over an authorised SignalR hub), and (iii) a configurable early-warning engine: every incident is scored 60% type-severity + 40% spatio-temporal frequency against administrator-editable RiskRules, raising de-duplicated ConflictAlerts on density or critical-severity triggers. The model centres on UserLocation → IncidentReport (+IncidentType, RiskScore/Level, assignment, resolution), ConflictAlert, RiskRule, AdministrativeArea (Benue/Makurdi/Guma/Logo seed), Evidence, StatusHistory and AuditLog over Identity tables. Evidence examined: Release folder-publish, 665-line InitialCreate migration (7 types, default 3-in-5 km-in-24 h rule), and a 3,027-line boot/migration log. No automated tests, Dockerfile or CI were found; verification is manual + log evidence. The report gives deployment/ops procedures, a STRIDE threat model with hardening steps, and limitations/future work.').runs[0].font.size = Pt(10)

# ---------- 1 ----------
H1('1  Introduction')
H2('1.1  Background and objectives')
doc.add_paragraph('Competition over land, water and grazing routes in Benue State (Makurdi, Guma, Logo — the seeded default areas) can escalate from disputes and livestock intrusion to destruction, threats and physical violence. Word-of-mouth reporting, imprecise locations and no shared live picture delay response. The project builds a low-friction reporting channel plus a GIS operations picture with automatic escalation when incidents cluster in space and time.').runs[0].font.size = Pt(10)
B('public reporting without login: consented GPS share on a mobile-friendly /Report page, validated and rate-limited server-side.', 'Frictionless reporting — ')
B('live common picture: dashboard (stats, alerts, live table, Leaflet map), filterable worklist, GeoJSON ops map — all SignalR-pushed.', 'Situational awareness — ')
B('explainable early warning: RiskAssessmentService + AlertService driven by editable RiskRules, never hard-coded constants.', 'Automatic escalation — ')
B('accountability: role-based access, per-incident status workflow with history, global audit log on MySQL with seeded taxonomy.', 'Audit — ')
H2('1.2  Scope and stakeholders')
T(['Stakeholder', 'Use'], [
    ['Community member (anonymous)', 'Share GPS + report; receives reference (e.g. INC-00025), risk, nearest area.'],
    ['Monitoring / Response Officer', 'Watch dashboard/map, triage, assign, advance status, acknowledge alerts.'],
    ['Administrator / SuperAdmin', 'Officer powers + user roles + risk-threshold tuning.'],
    ['Viewer', 'Read-only authorised views.'],
])
P('Out of scope: predictive modelling, SMS/USSD ingestion, native apps, automated dispatch — the system stops at detection, visualisation and audited human workflow.', italic=True, size=9)

# ---------- 2 ----------
H1('2  Requirements and Threat Model')
H2('2.1  Functional and non-functional requirements')
T(['ID', 'Requirement → realised in'], [
    ['FR-1', 'Anonymous consented GPS/incident report → /Report, gps.js, LocationShareDto, LocationService'],
    ['FR-2/3', 'Taxonomy with severity weights; auto score 0–100 + Low/Med/High/Critical → IncidentType, RiskAssessmentService'],
    ['FR-4', 'Hotspot/severity alerts with de-dupe → AlertService, ConflictAlert'],
    ['FR-5–8', 'Live dashboard + worklist + detail triage + GeoJSON ops map → Dashboard/Incidents/Map pages + live-map/monitoring.js'],
    ['FR-9–12', 'Alert ack; configurable rules; user admin; dual audit → Admin pages; AuditLog + StatusHistory'],
])
B('Bootstrap 5 + one-tap GPS with retry; Leaflet/OSM tiles; toasts for live events.', 'Usability: ')
B('Bounding-box pre-filter; caps (30 recent / 10 alerts / 200 list / 500 GeoJSON); 2 h session + memory cache.', 'Performance: ')
B('Auto-migrate + seed (Program.cs); Serilog console + daily file; handler + HSTS outside Development.', 'Reliability: ')
B('net8.0 + Pomelo 8.0.2 (AutoDetect, 8.0.36 fallback); folder publish Release/Any CPU.', 'Portability: ')
H2('2.2  Trust boundaries, roles, STRIDE')
P('Boundaries: (B1) browser↔server — GPS/clocks/inputs untrusted (validate, rate-limit, consent flag, separate ReceivedAt); (B2) public (/Report, Login) ↔ authorised (/Dashboard, /Incidents, /Map, /Admin, hub) via AuthorizeFolder + [Authorize]/roles; (B3) app ↔ MySQL credentials and ↔ OSM/Nominatim/CDN third parties. Roles (Enums.cs): SuperAdmin, Administrator (full incl. users/rules); MonitoringOfficer, ResponseOfficer (monitor/triage/ack); Viewer (read-only). Login /Account/Login, denied /Account/AccessDenied; light password policy (digit + length 6, no lockout/2FA) — see §7.').runs[0].font.size = Pt(10)
T(['Threat', 'Example → mitigation / residual'], [
    ['Spoofing', 'Fake GPS flood; stolen creds → Identity+hub auth; 10/min limit; audit / RES: weak passwords, empty dev DB secret'],
    ['Tampering', 'Altered coords/params → range checks; antiforgery (forms + fetch header) / RES: coarse officer powers'],
    ['Repudiation', '“Not my status change” → AuditLog + StatusHistory (who/IP/when/comment)'],
    ['Disclosure', 'Reporter PII on live map; conn-string in config → closed zone / RES: minimise DTOs; vault secrets'],
    ['DoS', 'Report flood; heavy scans → rate caps + query caps + prefilter / RES: single-node limiter'],
    ['Elevation', 'Viewer→Admin CSRF → role attributes + hidden nav + antiforgery / RES: add fine-grained policies'],
])
P('Handled abuse cases: Consent must be true; out-of-range fixes rejected; alerts suppressed when covered by an unacknowledged alert in the same window/radius; Critical severity can alert below the density threshold.', italic=True, size=9)

# ---------- 3 ----------
H1('3  System Design')
H2('3.1  Architecture and flows')
P('Three-tier monolith (Plan §§9,12): Browser (Bootstrap/jQuery/DataTables, Leaflet 1.9.4, SignalR JS 8.0.0, Geolocation) → ASP.NET Core 8 (page handlers, scoped services, MonitoringHub, Identity, EF Core, Serilog) → MySQL 8 (FarmerHerderConflictDb). No separate Web API: POST /Report?handler=Submit, GET /Map?handler=GeoJson, Assign/Status/Ack/Role handlers, plus hub /monitoringHub (Program.cs maps only MapRazorPages + MapHub). Runtime third parties: OSM tiles + Nominatim (client-side).').runs[0].font.size = Pt(10)
C('Browser (/Report,/Dashboard,/Map) ── fetch + SignalR ──▶ ASP.NET Core 8 (Pages+Services+Hub+Identity)\n                                                 │ EF Core/Pomelo              │ SignalR “dashboard”\n                                                 ▼                             ▼\n                                    MySQL FarmerHerderConflictDb      Officer browsers (tables/maps/toasts)', title='Figure 3-1 — Logical flow.')
B('gps.js (consent → high-accuracy fix, 15 s → Nominatim → POST Submit) → SaveAsync → nearest-area → CreateAsync (default “Other”) → EvaluateAsync → CheckAndGenerateAsync → broadcasts ReceiveIncident/ReceiveAlert → JSON {ref, risk, alerts, place}.', 'Report: ')
B('Dashboard seeds 30 recent + 10 alerts into initLiveMap [7.7322,8.5391]×10; monitoring.js joins “dashboard”; ReceiveIncident adds marker/row/toast, ReceiveAlert raises banner. Hub only joins/leaves.', 'Live: ')
B('/Incidents (200 + DataTables) → Details/{id} (map z14, assign + status/comment, history) → /Map?focus=lat,lon; /Admin/Alerts ack; /Map GeoJSON {kind: incident|alert|hotspot}, hotspot L.circle(r×1000), colours Critical red/High orange/Medium gold/Low green.', 'Triage/GIS: ')
H2('3.2  Data and early-warning design')
T(['Entity', 'Keys / notes'], [
    ['UserLocation → IncidentReport', '1—* (Restrict); fix has SessionId, lat/lon, acc, CapturedAt (browser) vs ReceivedAt (server), IP'],
    ['IncidentReport (+Type, +Officer)', 'UQ ReferenceNumber INC-#####; idx (lat,lon)/ReportedAt/Status; RiskScore/Level; Status New→…→Resolved; reporter/community/LGA; Cascade history/evidence'],
    ['ConflictAlert / RiskRule', 'Alert: title/msg, lat/lon, radius 5, Trigger Frequency|Severity, count, ack fields. Rule: min 3, 5 km, 24 h, High 50, Critical 75'],
    ['Area / Evidence / History / Audit', 'Area self-parent (Benue+Makurdi/Guma/Logo; GeoJSON opt). Evidence modelled, no UI. History + AuditLog (actor/IP/time)'],
    ['Seeds', '7 types: Suspicious 20/Low, Dispute 40, Intrusion 35, Threat 60, Destruction 65, Physical Conflict 90/Critical, Other 15'],
])
P('Scoring (RiskAssessmentService): freq=min(100, nearby×100/minCount); score=round(0.6×severity+0.4×freq), 0–100; level vs rule High/Critical with Medium floor 25 (≈ Plan bands 0–24/25–49/50–74/75–100). E.g. Threat(60)+2 neighbours (min 3): freq 66 → 62 High; Physical Conflict(90)+3: 94 Critical. Hotspots: greedy clustering of recent unresolved incidents (seed→neighbourhood ≥ minCount → centroid + max-score level). Alerts iterate active rules, skip overlap with unacknowledged alerts in window/radius, raise HIGH-RISK HOTSPOT (Frequency) or CRITICAL incident (Severity).').runs[0].font.size = Pt(10)

# ---------- 4 ----------
H1('4  Implementation')
T(['Layer', 'As-built choice (evidence)'], [
    ['Runtime/UI', 'ASP.NET Core 8 net8.0, Nullable+ImplicitUsings; Razor Pages (Program.cs:42–48,94)'],
    ['ORM/DB/Auth', 'EF Core 8.0.8 + Pomelo 8.0.2 (UseMySql, AutoDetect/8.0.36); Identity.EF 8.0.8, cookie Login/Denied paths'],
    ['Realtime/GIS', 'SignalR server + @microsoft/signalr 8.0.0; Leaflet 1.9.4 + OSM; Geolocation + Nominatim; DataTables 1.13.8'],
    ['Logging/state', 'Serilog.AspNetCore 8.0.1 (console + Logs/app-.log daily); memory cache + 2 h session'],
])
P('Pipeline (Program.cs): Serilog → handler/HSTS (non-dev) → HTTPS → static → routing → session → auth → pages + hub. Startup runs MigrateAsync + SeedAsync (roles, admin@test.com, 7 types, rule, 4 areas) with logged non-fatal failure. Services (all Scoped): IGisService, IRiskAssessmentService, ILocationService, IIncidentService, IAlertService, IAuditService, INotificationService. Schema: migration 665 lines (utf8mb4; Identity + domain tables); DbContext sets ReferenceNumber UQ, spatial/reporting indexes, Restrict parents, Cascade children, self-FK areas.').runs[0].font.size = Pt(10)
C('Score = Round(sev*0.6 + Min(100, nearby*100/minCount)*0.4); Classify: ≥crit→Critical; ≥high→High; ≥25→Medium else Low. // RiskAssessmentService.cs:21–46\nnearby = GetNearbyAsync(lat,lon,radius,hours): box [±r/111, ±r/(111*Max(.2,cos φ))] × ReportedAt≥since → Haversine filter (EarthKm 6371). // GisService.cs:20–49\nAlert: skip if covered by unacked alert in window/radius; fire if count≥min OR Critical → {HIGH-RISK HOTSPOT|CRITICAL incident, Frequency|Severity}. // AlertService.cs\nINC-{(Max(Id)+1):D5}; broadcast hub.Group("dashboard").SendAsync("ReceiveIncident"/"ReceiveAlert"). // IncidentService, NotificationService, [Authorize] Hub',
  title='Excerpts 4-A — Scoring, GIS prefilter, alert de-dupe, reference + push (shortened).')
P('Pages: / (hero + CTAs + 3 cards); /Report public (consent, Share/Retry, lat/lon/acc/place; gps.js only — pure shares use “Other”; richer IncidentSubmitDto exists); Login (ReturnUrl ~/Dashboard, seed hint) / Logout→/Index; Dashboard (● LIVE, 5 cards #statTotal/Active/Today/Critical/Unresolved, #alertBanner, 420 px map, alert + 30-row tables, JSON seed + connectMonitoring); Incidents (q/status/risk/type, DataTables); Details/{id} (badges, /Map?focus link, detail list, 320 px map z14, assign + status/comment, history); Map (risk/status/date/hotspots, 600 px map, GeoJSON, focus z15); Admin/Alerts (ack + map links); RiskRules (list + Edit all thresholds, audited); Users (single-role replace); _Layout (Conflict Monitor brand, role-gated nav). DTOs/Dtos.cs: Location/Submit/Share (ranges, Consent=true) + LiveLocation/Alert/Hotspot read models; layered DataAnnotations + Validate + rate gate + antiforgery; incident lat/lon copied from stored fix.').runs[0].font.size = Pt(10)

# ---------- 5 ----------
H1('5  Testing and Evaluation')
P('No test project/framework, Dockerfile, CI or README found — solution holds only the web project; docs are the Plan + comments. Evidence: 3,027-line log (cold-start empty-server errors → CREATE DATABASE → __EFMigrationsHistory → Applying 20261008171902_InitialCreate → CREATE TABLEs) and prior Release publish (bin/Release, obj/Release/PubTmp/Out). Below is the repeatable manual matrix; evaluation follows.').runs[0].font.size = Pt(10)
T(['#', 'Case → expected (as-built)'], [
    ['T-1–3', 'Valid consented share → 200 {INC-#####, risk, area} + rows + push (alert iff rule hit). No-consent/bad lat-lon → 400, nothing saved. 11th share/60 s → 429.'],
    ['T-4–6', 'Anon /Dashboard,/Map,/Incidents,/Admin, hub join → Login redirect / hub reject. admin@test.com/Admin123! → cookie + admin nav. Viewer → /Admin/* → 403/AccessDenied.'],
    ['T-7–9', '3 reports ≤5 km/24 h → HIGH-RISK HOTSPOT (High). Lone Physical Conflict → High/Critical (Severity path). Repeat inside unacked window → suppressed.'],
    ['T-10–13', 'Ack/assign/status → ack flag + history + audit; ResolvedAt; stats drop. GeoJSON filters → kind-tagged Points + circles; ?focus centres. Rule edit → persisted + audited + live. Cold empty-MySQL boot → migrate + seed.'],
])
P('Coverage: FR-1–FR-12 happy paths complete; pilot-grade NFRs (caps, prefilter, rotation) adequate but unmeasured — no load/soak/accuracy/offline tests recorded. Main gap is calibration: 60/40 blend and defaults (3/5 km/24 h/50/75) are plausible, unevidenced — back-test on dated history + sensitivity analysis before freezing thresholds. Known defects: no regression safety for scoring/antimeridian/authz; GPS-share fixed to “Other”; Evidence unwired; single-node limiter.').runs[0].font.size = Pt(10)

# ---------- 6 ----------
H1('6  Deployment and Operating Policy')
T(['Need', 'As-built / action'], [
    ['Runtime/DB', '.NET 8 runtime; MySQL 8 (≈8.0.36); MigrateAsync creates FarmerHerderConflictDb if absent'],
    ['Config', 'DefaultConnection localhost:3306/root (empty dev password; code fallback root) — override by env in prod; rotate SeedAdmin now'],
    ['Ports/net', 'Dev 5278/7006, IISEx 32583/44314; prod behind TLS reverse proxy; egress to OSM/Nominatim/CDN unless vendored'],
])
C('dotnet restore; dotnet tool install --global dotnet-ef   # once\ndotnet ef database update --project Assignment/Assignment.csproj\ndotnet run --project Assignment/Assignment.csproj --launch-profile https  # https://localhost:7006 (admin@test.com / Admin123!)',
  title='Routine 6-A — Restore, migrate, run.')
P('Publish: FolderProfile.pubxml (Release/Any CPU, FileSystem → C:\\Users\\BDIC-004\\Desktop\\MyAssignment, last 2026-10-08T17:59:47). Policy: publish Release, set Production env, inject secrets via environment (never ship real ones), front with IIS/Nginx + TLS, schedule MySQL dumps + Logs retention; boot health = admin login + Dashboard stats + Map GeoJSON features. Shifts: keep Dashboard open (● LIVE); verify before ack (acked radius masks new clusters); tune one rule variable at a time (audited); single-role user changes; weekly audit review; verify Community/LGA spellings; treat ReceivedAt authoritative; back up before seed/rule changes.').runs[0].font.size = Pt(10)

# ---------- 7 ----------
H1('7  Security Considerations')
B('Closed-by-default zone (4× AuthorizeFolder + page/hub attributes); tiny public surface; server-side GPS (ReceivedAt, ranges, consent, 10/min); antiforgery on forms + JSON header; 5-role least privilege + hidden admin nav + single-role replace; dual audit + Serilog; HTTPS/HSTS + parameterised EF queries.', 'Strengths: ')
T(['#', 'Weakness → fix'], [
    ['S-1', 'Seed admin@test.com/Admin123! + hint; empty/fallback DB password (High) → rotate, vault/env secrets, fail closed, drop hint'],
    ['S-2', 'Weak policy (6+digit) + no lockout/2FA (High) → ≥12/complexity, lockout + 2FA for admins, failure alerts'],
    ['S-3', 'Reporter PII stored + broadcast (Med) → minimise/mask by role, retention/purge, consent notice'],
    ['S-4', 'Coarse powers — any officer acks/assigns (Med) → policies (e.g. ack ≥ ResponseOfficer, team/LGA scope) + T-4/T-6 tests'],
    ['S-5/6', 'Local-only limiter; CDN/Nominatim w/o SRI (Med/Low) → distributed throttle + CAPTCHA; vendor/pin + SRI + offline fallback'],
])
P('Transport/secrets/logging: TLS 1.2+ end to end, HSTS preload, HttpOnly+Secure+SameSite cookies; never commit real secrets (current strings are dev placeholders); Information/Warning log levels; no PII/secrets in logs; ≥90-day secured retention; least-privilege MySQL + tested encrypted backups.').runs[0].font.size = Pt(10)

# ---------- 8 + 9 ----------
H1('8  Limitations and Future Work')
B('No tests/CI/containers; naive geo (doubles + app Haversine, B-tree indexes only, O(n²) greedy hotspots); single-node session/limiter/hub; GPS-share fixed “Other”, evidence unwired; linear area scan, hard-coded 25 floor, uncalibrated defaults; offline gaps (tiles/queue/SMS); English-only, low-bandwidth/a11y gaps.', 'Limits: ')
T(['Pri', 'Item → acceptance'], [
    ['P0', 'xUnit + WebApplicationFactory (scoring, poles/antimeridian, dedupe, authz) + CI gate → green T-4/T-6 per PR'],
    ['P0', 'Vault/env secrets, strong policy, lockout+2FA → seed creds dead after boot; brute-force blocked'],
    ['P1', 'SPATIAL/PostGIS + NetTopologySuite, DB-side distance, paged GeoJSON → p95 <500 ms @100k'],
    ['P1', 'Expose type/description + evidence upload (AV/size/type) on /Report → renders on Details'],
    ['P1/P2', 'Redis backplane + distributed throttle + SRI/vendor maps; calibrate on history (ROC); offline pack + SMS queue'],
    ['P3', 'Chart.js trends, LGA tables, CSV/GeoJSON export, boundary editor → one-click monthly report'],
])
H1('9  Conclusion')
doc.add_paragraph('A consented GPS share becomes a scored, mapped, auditable incident within a minute, with automatic escalation on clustering or grave severity. The build is faithful to its plan, the model is pilot-ready for Benue State, and the warning maths are transparent and tunable. Trust it for a controlled pilot; before wider use, close the P0 items — tests + CI, secret/auth hardening, threshold calibration — then pursue the spatial, resilience and analytics increments. That order yields a sound base for data-driven prevention.').runs[0].font.size = Pt(10)

# ---------- REFS ----------
H1('References')
for i, r_ in enumerate([
    'Microsoft. ASP.NET Core 8 — Razor Pages, Identity, SignalR, EF Core. https://learn.microsoft.com/aspnet/core (Oct 2026).',
    'Pomelo 8.0.2 (ServerVersion.AutoDetect); Serilog.AspNetCore 8.0.1; Leaflet 1.9.4 + OSM/Nominatim policies; SignalR JS 8.0.0; DataTables 1.13.8; Bootstrap 5.',
    'As-built sources: Program.cs; appsettings.json; ApplicationDbContext.cs; DbInitializer.cs; 20261008171902_InitialCreate; Models/*.cs; Dtos.cs; Services/*.cs; MonitoringHub.cs; Pages/**; gps.js; live-map.js; monitoring.js; app-20261008.log; launchSettings.json; FolderProfile.pubxml; Implementation_Plan.md.',
    'OWASP Top 10 (2021) + ASVS — basis for §7. https://owasp.org.',
], 1):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1)
    a = p.add_run(f'[{i}]  '); a.bold = True; a.font.size = Pt(8.5)
    b = p.add_run(r_); b.font.size = Pt(8.5)

H1('Appendix A: Command Reference')
P('From solution root (PowerShell 5.1).', italic=True, size=8.5)
T(['Purpose', 'Command'], [
    ['Restore / run', 'dotnet restore · dotnet run --project Assignment\\Assignment.csproj --launch-profile https (7006)'],
    ['EF Core', 'dotnet tool install --global dotnet-ef · ef migrations add <N> / database update / migrations script -o migration.sql'],
    ['Publish', 'dotnet publish Assignment\\Assignment.csproj -c Release -o C:\\Users\\BDIC-004\\Desktop\\MyAssignment'],
    ['Logs/smoke', 'Get-Content Assignment\\Logs\\app-20261008.log -Tail 50 -Wait · Invoke-RestMethod ".../Map?handler=GeoJson&showHotspots=true"'],
    ['MySQL', 'CREATE DATABASE FarmerHerderConflictDb CHARACTER SET utf8mb4; mysqldump -u root -p FarmerHerderConflictDb > backup.sql'],
])
C('{"ConnectionStrings":{"DefaultConnection":"server=localhost;port=3306;database=FarmerHerderConflictDb;user=root;password=;"},\n "Serilog":{"MinimumLevel":"Information","WriteTo":[{"Name":"Console"},{"Name":"File","Args":{"path":"Logs/app-.log","rollingInterval":"Day"}}]},\n "SeedAdmin":{"Email":"admin@test.com","Password":"Admin123!"}}', title='Config — appsettings.json (dev placeholder; override in prod).')

H1('Appendix B: Source Code Layout')
P('Excludes bin/obj/.vs/wwwroot/lib. JS line counts exact.', italic=True, size=8.5)
C('AssignmentProject.slnx → Assignment/Assignment.csproj (net8.0; Identity.EF/Design 8.0.8; Pomelo 8.0.2; Serilog 8.0.1)\nProgram.cs (pipeline+AuthZ+DI+Migrate/Seed+Pages/Hub) · appsettings*.json · Plan.md\nData/: ApplicationDbContext (9 sets+indexes+seeds) · DbInitializer (5 roles+admin) · Migrations/20261008171902_InitialCreate (665)\nDTOs/Dtos.cs · Hubs/MonitoringHub ([Authorize], dashboard group) · Models/ (User,Location,Type,Report,Alert,Rule,Area,Evidence,History,Audit,Enums)\nServices/: Gis (Haversine/box/hotspot/GeoJSON) · Risk (60/40) · Alert (+Audit) · Incident (INC-#####/stats/status/assign) · Location (10/min) · Notify\nPages/: Index|Privacy|Error · Account/Login|Logout|Denied · Report (public Submit) · Dashboard (live) · Incidents/Index|Details · Map (GeoJSON) · Admin/Alerts|RiskRules|Users · Shared/_Layout\nwwwroot/js: gps.js(117) · live-map.js(69) [7.7322,8.5391] · monitoring.js(51) · site.js/css stock\nLogs/app-20261008.log (3027) · launchSettings (5278/7006) · FolderProfile.pubxml → Desktop\\MyAssignment',
  title='Layout B-1 — Filtered tree.')
P('End of report (10-page edition).', italic=True, size=9)

doc.save(OUT)
print(f'Saved to {OUT}')
